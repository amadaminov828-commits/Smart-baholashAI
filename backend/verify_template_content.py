import os
import django
import sys
import docx

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'valuation_platform.settings')
django.setup()

from reports.models import ReportTemplate

def analyze_template_content():
    templates = ReportTemplate.objects.all().order_by('-created_at')[:3]
    with open('template_dump.txt', 'w', encoding='utf-8') as f:
        f.write("Analyzing the last 3 templates...\n\n")
        
        for template in templates:
            f.write(f"--- Template: {template.name} (ID: {template.id}) ---\n")
            f.write(f"Path: {template.file.path}\n")
            
            try:
                doc = docx.Document(template.file.path)
                tags_found = []
                text_preview = ""
                
                common_tags = ['{raqami}', '{davlat_raqami}', '{modeli}', '{avto_modeli}', '{f_i_sh}', '{mulk_egasi}', '{shassi_raqami}', '{kuzov_raqami}']
                
                paragraphs = list(doc.paragraphs)
                for p in paragraphs:
                    text_preview += p.text + "\n"
                    for tag in common_tags:
                        if tag in p.text.lower():
                            tags_found.append(tag)
                
                f.write(f"Tags found: {set(tags_found)}\n")
                
                text_lower = text_preview.lower()
                if 'chery' in text_lower or 'tiggo' in text_lower:
                    f.write("WARNING: Found 'Chery' or 'Tiggo' hardcoded in text. Is this a generated report used as a template?\n")
                    
                if '01a123aa' in text_lower or '01 a 123 aa' in text_lower:
                    f.write("WARNING: Found standard plate '01 A 123 AA' hardcoded.\n")
                    
                if len(set(tags_found)) < 2:
                    f.write("WARNING: Very few standard `{tags}` found. This might be a generated report, not a template with placeholders.\n")
                    f.write(f"Sample text snippet:\n{text_preview[:500]}\n")
                    
            except Exception as e:
                f.write(f"Error reading docx: {e}\n")
            f.write("\n")

if __name__ == '__main__':
    analyze_template_content()
